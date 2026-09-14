
#!/usr/bin/env python3
"""
PolicyGuard AI - Production Multimodal Document Parser
=======================================================

Unified parser for:
- PDF
- DOCX
- XLSX
- PPTX
- Images (PNG/JPG/JPEG/TIFF)

Features:
- OCR fallback for scanned PDFs/images
- Sentence-aware chunking
- Batch processing with thread pooling
- Stable input ordering for batch results
- File metadata extraction
- Thread-safe global singleton
- Defensive dependency handling
- Windows/Linux path compatibility

Important:
- .doc and .xls are NOT treated as DOCX/XLSX internally.
  They are rejected unless a compatible backend is added.
"""

import hashlib
import logging
import os
import re
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple, Union

# ---------------------------------------------------------------------------
# Project path setup
# ---------------------------------------------------------------------------

current_file = Path(__file__).resolve()
project_root = current_file.parent.parent.parent

if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

from config.settings import settings
from src.core.exceptions import RAGException, ValidationError

logger = logging.getLogger(__name__)


# =============================================================================
# CONSTANTS
# =============================================================================

SUPPORTED_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tiff", ".tif"}

# These extensions are intentionally NOT mapped to DOCX/XLSX.
# python-docx does not parse legacy .doc files and openpyxl does not parse
# legacy .xls files.
SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".pptx",
    *SUPPORTED_IMAGE_EXTENSIONS,
}

DEFAULT_MAX_WORKERS = 4
MIN_CHUNK_SIZE = 1
MAX_CHUNK_SIZE = 1_000_000

DEFAULT_MAX_FILE_SIZE_BYTES = 50 * 1024 * 1024
MAX_FILENAME_LENGTH = 255
MAX_TEXT_LENGTH = 10_000_000
VALID_NAMESPACES = {"policy", "talent"}
_SCOPE_PATTERN = re.compile(r"^[A-Za-z0-9_.:@-]{1,128}$")


# =============================================================================
# UTILITY FUNCTIONS
# =============================================================================

def _utc_now_iso() -> str:
    """Return the current UTC time as an ISO-8601 string."""
    return datetime.now(timezone.utc).isoformat()


def _safe_int(value: Any, default: int) -> int:
    """Convert a value to int without allowing malformed configuration to fail."""
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _get_default_chunk_size() -> int:
    """Read chunk size from settings with a safe fallback."""
    return _safe_int(getattr(settings, "CHUNK_SIZE", 300), 300)


def _get_default_chunk_overlap() -> int:
    """Read chunk overlap from settings with a safe fallback."""
    return _safe_int(getattr(settings, "CHUNK_OVERLAP", 50), 50)


def _validate_chunk_config(chunk_size: int, chunk_overlap: int) -> Tuple[int, int]:
    """
    Validate and normalize chunking configuration.

    Overlap must always be smaller than chunk size. Otherwise the parser can
    repeatedly revisit the same text region.
    """
    if not isinstance(chunk_size, int) or isinstance(chunk_size, bool):
        raise ValidationError("chunk_size must be an integer")

    if not isinstance(chunk_overlap, int) or isinstance(chunk_overlap, bool):
        raise ValidationError("chunk_overlap must be an integer")

    if not MIN_CHUNK_SIZE <= chunk_size <= MAX_CHUNK_SIZE:
        raise ValidationError(
            f"chunk_size must be between {MIN_CHUNK_SIZE} and {MAX_CHUNK_SIZE}"
        )

    if chunk_overlap < 0:
        raise ValidationError("chunk_overlap cannot be negative")

    if chunk_overlap >= chunk_size:
        raise ValidationError(
            "chunk_overlap must be smaller than chunk_size"
        )

    return chunk_size, chunk_overlap


def _normalize_path(file_path: Union[str, Path]) -> Path:
    """Convert a user-supplied path to a normalized Path object."""
    if isinstance(file_path, Path):
        return file_path

    if not isinstance(file_path, str):
        raise ValidationError("file_path must be a string or pathlib.Path")

    if not file_path.strip():
        raise ValidationError("file_path cannot be empty")
    if "\x00" in file_path:
        raise ValidationError("file_path contains a null byte")

    return Path(file_path).expanduser()



def _get_max_file_size_bytes() -> int:
    """Read the configured maximum upload size with a safe fallback."""
    raw = getattr(settings, "MAX_UPLOAD_SIZE_MB", None)
    if raw is None:
        raw = getattr(settings, "MAX_FILE_SIZE_MB", None)
    try:
        mb = int(raw)
        if mb > 0:
            return mb * 1024 * 1024
    except (TypeError, ValueError):
        pass
    return DEFAULT_MAX_FILE_SIZE_BYTES


def _validate_scope(value: Optional[str], field_name: str, default: str) -> str:
    """Validate organization/namespace values before they enter metadata."""
    resolved = default if value is None else str(value).strip()
    if not resolved or not _SCOPE_PATTERN.fullmatch(resolved):
        raise ValidationError(f"Invalid {field_name}")
    return resolved


def _safe_filename(name: str) -> str:
    """Return a bounded filename without path components."""
    safe = Path(str(name)).name
    if not safe or safe in {".", ".."} or len(safe) > MAX_FILENAME_LENGTH:
        raise ValidationError("Invalid filename")
    return safe


def _file_sha256(path: Path) -> str:
    """Hash the source file for stable document identity."""
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()

# =============================================================================
# CHUNKING UTILITIES
# =============================================================================

class TextChunker:
    """
    Configurable sentence-aware text chunker.

    Chunk sizes are measured in characters rather than tokens so that the
    parser remains independent of any particular embedding/LLM tokenizer.
    """

    def __init__(
        self,
        chunk_size: Optional[int] = None,
        chunk_overlap: Optional[int] = None,
        min_chunk_size: int = 50,
        respect_sentences: bool = True,
    ):
        configured_size = (
            _get_default_chunk_size()
            if chunk_size is None
            else chunk_size
        )

        configured_overlap = (
            _get_default_chunk_overlap()
            if chunk_overlap is None
            else chunk_overlap
        )

        configured_size, configured_overlap = _validate_chunk_config(
            configured_size,
            configured_overlap,
        )

        if not isinstance(min_chunk_size, int) or isinstance(min_chunk_size, bool):
            raise ValidationError("min_chunk_size must be an integer")

        if min_chunk_size < 1:
            raise ValidationError("min_chunk_size must be greater than zero")

        if min_chunk_size > configured_size:
            min_chunk_size = configured_size

        self.chunk_size = configured_size
        self.chunk_overlap = configured_overlap
        self.min_chunk_size = min_chunk_size
        self.respect_sentences = bool(respect_sentences)

        logger.debug(
            "TextChunker initialized: size=%s overlap=%s "
            "min_size=%s sentences=%s",
            self.chunk_size,
            self.chunk_overlap,
            self.min_chunk_size,
            self.respect_sentences,
        )

    def chunk(
        self,
        text: str,
        source: str,
        metadata: Optional[Dict[str, Any]] = None,
    ) -> List[Dict[str, Any]]:
        """
        Split text into overlapping chunks.

        Returns:
            List of dictionaries containing:
            - content
            - metadata
            - score
        """
        if not isinstance(text, str):
            raise ValidationError("text must be a string")

        if not text.strip():
            return []

        source = str(source)
        base_metadata = dict(metadata or {})
        chunks: List[Dict[str, Any]] = []

        text_length = len(text)
        start = 0
        chunk_id = 0

        while start < text_length:
            original_start = start
            end = min(start + self.chunk_size, text_length)

            if self.respect_sentences and end < text_length:
                boundary = self._find_sentence_boundary(
                    text,
                    start,
                    end,
                )

                # Only accept a sentence boundary if it doesn't make the
                # chunk unreasonably small.
                if boundary > start + self.min_chunk_size:
                    end = boundary

            raw_chunk = text[start:end]
            chunk_text = raw_chunk.strip()

            if chunk_text:
                leading_trim = len(raw_chunk) - len(raw_chunk.lstrip())
                trailing_trim = len(raw_chunk) - len(raw_chunk.rstrip())

                char_start = start + leading_trim
                char_end = max(char_start, end - trailing_trim)

                # If this is a tiny tail, merge it into the previous chunk
                # rather than creating a useless fragment.
                if (
                    len(chunk_text) < self.min_chunk_size
                    and chunks
                ):
                    previous = chunks[-1]

                    separator = " " if previous["content"] else ""
                    previous["content"] = (
                        previous["content"] + separator + chunk_text
                    )

                    previous["metadata"]["char_end"] = char_end
                    previous["metadata"]["char_length"] = len(
                        previous["content"]
                    )

                    # Tail has been consumed completely.
                    start = text_length
                    break

                chunk_metadata = {
                    "source": source,
                    "chunk_id": chunk_id,
                    "char_start": char_start,
                    "char_end": char_end,
                    "char_length": len(chunk_text),
                    **base_metadata,
                }

                chunks.append(
                    {
                        "content": chunk_text,
                        "metadata": chunk_metadata,
                        "score": 0.0,
                    }
                )

                chunk_id += 1

            # Guarantee forward progress.
            next_start = end - self.chunk_overlap

            if next_start <= original_start:
                next_start = original_start + 1

            if next_start >= text_length:
                break

            start = next_start

        average_size = (
            sum(
                chunk["metadata"]["char_length"]
                for chunk in chunks
            )
            / max(len(chunks), 1)
        )

        logger.debug(
            "Chunked %s chars from '%s' -> %s chunks "
            "(avg %.0f chars/chunk)",
            text_length,
            source,
            len(chunks),
            average_size,
        )

        return chunks

    def _find_sentence_boundary(
        self,
        text: str,
        start: int,
        end: int,
    ) -> int:
        """
        Find a reasonable sentence boundary near the end of a chunk.

        Supports:
        - period
        - exclamation mark
        - question mark
        - common whitespace after punctuation

        If no suitable boundary is found, the original end is returned.
        """
        search_text = text[start:end]

        # Search up to 150 characters backwards from the target boundary.
        min_index = max(0, len(search_text) - 150)

        for i in range(len(search_text) - 1, min_index - 1, -1):
            if search_text[i] not in ".!?":
                continue

            # Don't split decimal numbers such as "1.67".
            previous_char = (
                search_text[i - 1]
                if i > 0
                else ""
            )
            next_char = (
                search_text[i + 1]
                if i + 1 < len(search_text)
                else ""
            )

            if previous_char.isdigit() and next_char.isdigit():
                continue

            if i + 1 >= len(search_text):
                return start + i + 1

            if search_text[i + 1].isspace():
                return start + i + 1

        return end


# =============================================================================
# PARSER BACKENDS
# =============================================================================

class ParserBackend:
    """Base class for format-specific parser backends."""

    @staticmethod
    def is_available() -> bool:
        raise NotImplementedError

    @staticmethod
    def parse(file_path: Path) -> Dict[str, Any]:
        raise NotImplementedError


class PDFBackend(ParserBackend):
    """PDF parser using pypdf with OCR fallback."""

    @staticmethod
    def is_available() -> bool:
        try:
            import pypdf  # noqa: F401
            return True
        except ImportError:
            return False

    @staticmethod
    def parse(file_path: Path) -> Dict[str, Any]:
        result: Dict[str, Any] = {
            "text": "",
            "pages": 0,
            "has_images": False,
            "ocr_used": False,
            "parse_method": "pypdf",
        }

        try:
            from pypdf import PdfReader

            reader = PdfReader(str(file_path))
            result["pages"] = len(reader.pages)

            page_texts: List[str] = []

            for page_number, page in enumerate(reader.pages, start=1):
                try:
                    page_text = page.extract_text() or ""

                    if page_text.strip():
                        page_texts.append(
                            f"[Page {page_number}]\n"
                            f"{page_text.strip()}"
                        )

                    # Best-effort image detection.
                    try:
                        if getattr(page, "images", None):
                            result["has_images"] = True
                    except Exception:
                        pass

                except Exception as page_error:
                    logger.warning(
                        "Failed to extract text from PDF page %s: %s",
                        page_number,
                        page_error,
                    )

            result["text"] = "\n\n".join(page_texts)

            # Scanned PDFs frequently have little/no extractable text.
            if len(result["text"].strip()) < 200:
                logger.info(
                    "PDF appears scanned or contains limited text "
                    "(%s chars); attempting OCR",
                    len(result["text"].strip()),
                )
                result = PDFBackend._parse_with_ocr(
                    file_path,
                    result,
                )

            logger.info(
                "PDF parsed: %s pages, %s chars, OCR=%s",
                result["pages"],
                len(result["text"]),
                result["ocr_used"],
            )

        except ImportError:
            logger.warning(
                "pypdf is not installed; attempting OCR fallback"
            )
            result = PDFBackend._parse_with_ocr(
                file_path,
                result,
            )

        except Exception as error:
            logger.error(
                "PDF parse failed: %s",
                error,
                exc_info=True,
            )

            # Attempt OCR even if the PDF parser itself fails.
            ocr_result = PDFBackend._parse_with_ocr(
                file_path,
                result,
            )

            if not ocr_result.get("text"):
                ocr_result["error"] = str(error)

            result = ocr_result

        return result

    @staticmethod
    def _parse_with_ocr(
        file_path: Path,
        result: Dict[str, Any],
    ) -> Dict[str, Any]:
        """Attempt OCR extraction for scanned PDFs."""
        try:
            from src.vision.advanced_ocr import extract_text_from_pdf

            ocr_text = extract_text_from_pdf(file_path)

            # Support both string and dictionary return styles.
            if isinstance(ocr_text, dict):
                text = str(ocr_text.get("text", "") or "")
            else:
                text = str(ocr_text or "")

            if text.strip():
                result["text"] = text
                result["ocr_used"] = True
                result["parse_method"] = "ocr"

                logger.info(
                    "PDF OCR successful: %s chars",
                    len(text),
                )
            else:
                logger.warning(
                    "PDF OCR returned no text"
                )

        except ImportError:
            logger.warning(
                "OCR module is not available for PDF parsing"
            )

        except Exception as error:
            logger.error(
                "PDF OCR failed: %s",
                error,
                exc_info=True,
            )
            result["ocr_error"] = str(error)

        return result


class DOCXBackend(ParserBackend):
    """Microsoft Word DOCX parser."""

    @staticmethod
    def is_available() -> bool:
        try:
            from docx import Document  # noqa: F401
            return True
        except ImportError:
            return False

    @staticmethod
    def parse(file_path: Path) -> Dict[str, Any]:
        result: Dict[str, Any] = {
            "text": "",
            "paragraphs": 0,
            "has_tables": False,
            "tables_count": 0,
        }

        try:
            from docx import Document

            doc = Document(str(file_path))

            paragraphs: List[str] = []

            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()

                if text:
                    paragraphs.append(text)

            result["paragraphs"] = len(paragraphs)
            result["text"] = "\n\n".join(paragraphs)

            if doc.tables:
                result["has_tables"] = True
                result["tables_count"] = len(doc.tables)

                table_texts: List[str] = []

                for table_number, table in enumerate(
                    doc.tables,
                    start=1,
                ):
                    rows: List[str] = []

                    for row in table.rows:
                        cells = [
                            cell.text.strip()
                            for cell in row.cells
                            if cell.text.strip()
                        ]

                        if cells:
                            rows.append(" | ".join(cells))

                    if rows:
                        table_texts.append(
                            f"[Table {table_number}]\n"
                            + "\n".join(rows)
                        )

                if table_texts:
                    table_text = "\n\n".join(table_texts)

                    if result["text"]:
                        result["text"] += "\n\n"

                    result["text"] += table_text

            logger.info(
                "DOCX parsed: %s paragraphs, %s tables, %s chars",
                result["paragraphs"],
                result["tables_count"],
                len(result["text"]),
            )

        except ImportError:
            result["error"] = "python-docx is not installed"

        except Exception as error:
            logger.error(
                "DOCX parse error: %s",
                error,
                exc_info=True,
            )
            result["error"] = str(error)

        return result


class XLSXBackend(ParserBackend):
    """Excel XLSX parser."""

    @staticmethod
    def is_available() -> bool:
        try:
            from openpyxl import load_workbook  # noqa: F401
            return True
        except ImportError:
            return False

    @staticmethod
    def parse(file_path: Path) -> Dict[str, Any]:
        result: Dict[str, Any] = {
            "text": "",
            "sheets": 0,
            "rows": 0,
            "columns": 0,
        }

        workbook = None

        try:
            from openpyxl import load_workbook

            workbook = load_workbook(
                str(file_path),
                read_only=True,
                data_only=True,
            )

            result["sheets"] = len(workbook.sheetnames)

            sheet_texts: List[str] = []
            total_rows = 0
            max_columns = 0

            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                rows_data: List[str] = []

                for row in sheet.iter_rows(values_only=True):
                    values: List[str] = []

                    for cell in row:
                        if cell is None:
                            values.append("")
                            continue

                        value = str(cell).strip()
                        values.append(value)

                    if any(values):
                        rows_data.append(
                            " | ".join(
                                value
                                for value in values
                                if value
                            )
                        )

                        total_rows += 1
                        max_columns = max(
                            max_columns,
                            len(values),
                        )

                if rows_data:
                    sheet_texts.append(
                        f"[Sheet: {sheet_name}]\n"
                        + "\n".join(rows_data)
                    )

            result["rows"] = total_rows
            result["columns"] = max_columns
            result["text"] = "\n\n".join(sheet_texts)

            logger.info(
                "XLSX parsed: %s sheets, %s rows, %s chars",
                result["sheets"],
                result["rows"],
                len(result["text"]),
            )

        except ImportError:
            result["error"] = "openpyxl is not installed"

        except Exception as error:
            logger.error(
                "XLSX parse error: %s",
                error,
                exc_info=True,
            )
            result["error"] = str(error)

        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass

        return result


class PPTXBackend(ParserBackend):
    """PowerPoint PPTX parser."""

    @staticmethod
    def is_available() -> bool:
        try:
            from pptx import Presentation  # noqa: F401
            return True
        except ImportError:
            return False

    @staticmethod
    def parse(file_path: Path) -> Dict[str, Any]:
        result: Dict[str, Any] = {
            "text": "",
            "slides": 0,
            "shapes": 0,
        }

        try:
            from pptx import Presentation

            presentation = Presentation(str(file_path))
            result["slides"] = len(presentation.slides)

            slide_texts: List[str] = []
            total_shapes = 0

            for slide_number, slide in enumerate(
                presentation.slides,
                start=1,
            ):
                parts = [f"[Slide {slide_number}]"]

                for shape in slide.shapes:
                    shape_text = getattr(shape, "text", "")

                    if isinstance(shape_text, str) and shape_text.strip():
                        parts.append(shape_text.strip())
                        total_shapes += 1

                if len(parts) > 1:
                    slide_texts.append("\n".join(parts))

            result["shapes"] = total_shapes
            result["text"] = "\n\n".join(slide_texts)

            logger.info(
                "PPTX parsed: %s slides, %s text shapes, %s chars",
                result["slides"],
                result["shapes"],
                len(result["text"]),
            )

        except ImportError:
            result["error"] = "python-pptx is not installed"

        except Exception as error:
            logger.error(
                "PPTX parse error: %s",
                error,
                exc_info=True,
            )
            result["error"] = str(error)

        return result


class ImageBackend(ParserBackend):
    """Image parser using the project's OCR implementation."""

    @staticmethod
    def is_available() -> bool:
        try:
            from src.vision.advanced_ocr import extract_text_from_image  # noqa: F401
            return True
        except ImportError:
            return False

    @staticmethod
    def parse(file_path: Path) -> Dict[str, Any]:
        result: Dict[str, Any] = {
            "text": "",
            "ocr_confidence": 0.0,
            "language": "en",
            "ocr_used": True,
            "parse_method": "ocr",
        }

        try:
            from src.vision.advanced_ocr import extract_text_from_image

            ocr_result = extract_text_from_image(file_path)

            if isinstance(ocr_result, dict):
                result["text"] = str(
                    ocr_result.get("text", "") or ""
                )

                try:
                    result["ocr_confidence"] = float(
                        ocr_result.get("confidence", 0.0) or 0.0
                    )
                except (TypeError, ValueError):
                    result["ocr_confidence"] = 0.0

                result["language"] = str(
                    ocr_result.get("language", "en") or "en"
                )

            elif isinstance(ocr_result, str):
                result["text"] = ocr_result

            elif ocr_result is not None:
                result["text"] = str(ocr_result)

            if result["text"].strip():
                logger.info(
                    "Image OCR successful: %s chars, confidence=%.2f",
                    len(result["text"]),
                    result["ocr_confidence"],
                )
            else:
                logger.warning(
                    "Image OCR returned no text"
                )

        except ImportError:
            logger.warning(
                "OCR module is not available for image parsing"
            )
            result["error"] = "OCR not available"

        except Exception as error:
            logger.error(
                "Image OCR error: %s",
                error,
                exc_info=True,
            )
            result["error"] = str(error)

        return result


# =============================================================================
# MAIN MULTIMODAL PARSER
# =============================================================================

class MultimodalParser:
    """
    Unified document parser.

    Supported formats:
        PDF, DOCX, XLSX, PPTX, PNG, JPG, JPEG, TIFF, TIF
    """

    EXTENSION_BACKENDS = {
        ".pdf": PDFBackend,
        ".docx": DOCXBackend,
        ".xlsx": XLSXBackend,
        ".pptx": PPTXBackend,
        ".png": ImageBackend,
        ".jpg": ImageBackend,
        ".jpeg": ImageBackend,
        ".tiff": ImageBackend,
        ".tif": ImageBackend,
    }

    def __init__(
        self,
        chunk_size: Optional[int] = None,
        chunk_overlap: Optional[int] = None,
        max_workers: int = DEFAULT_MAX_WORKERS,
    ):
        if not isinstance(max_workers, int) or isinstance(max_workers, bool):
            raise ValidationError("max_workers must be an integer")

        if max_workers < 1 or max_workers > 64:
            raise ValidationError(
                "max_workers must be between 1 and 64"
            )

        self.chunker = TextChunker(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
        )

        self.max_workers = max_workers

        self.available_backends: Dict[str, type] = {}

        for extension, backend_class in self.EXTENSION_BACKENDS.items():
            try:
                available = backend_class.is_available()
            except Exception as error:
                available = False
                logger.warning(
                    "Backend availability check failed for %s: %s",
                    extension,
                    error,
                )

            if available:
                self.available_backends[extension] = backend_class
                logger.debug(
                    "Backend available: %s -> %s",
                    extension,
                    backend_class.__name__,
                )
            else:
                logger.debug(
                    "Backend unavailable: %s -> %s",
                    extension,
                    backend_class.__name__,
                )

        logger.info(
            "MultimodalParser initialized: %s backends available, "
            "chunk_size=%s, overlap=%s, workers=%s",
            len(self.available_backends),
            self.chunker.chunk_size,
            self.chunker.chunk_overlap,
            self.max_workers,
        )

    # -------------------------------------------------------------------------
    # Single-file parsing
    # -------------------------------------------------------------------------

    def parse_file(
        self,
        file_path: Union[str, Path],
        extract_metadata: bool = True,
        chunk: bool = True,
        *,
        organization_id: Optional[str] = None,
        namespace: str = "policy",
        document_id: Optional[str] = None,
    ) -> Dict[str, Any]:
        """
        Parse one supported document.

        Raises:
            FileNotFoundError
            ValidationError
            RAGException
        """
        path = _normalize_path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"File not found: {path}"
            )

        if not path.is_file():
            raise ValidationError(
                f"Path is not a file: {path}"
            )

        if path.is_symlink():
            raise ValidationError("Symbolic-link uploads are not allowed")

        try:
            file_size = path.stat().st_size
        except OSError as exc:
            raise ValidationError("Unable to inspect uploaded file") from exc

        if file_size <= 0:
            raise ValidationError("Uploaded file is empty")

        max_size = _get_max_file_size_bytes()
        if file_size > max_size:
            raise ValidationError(
                f"File exceeds the maximum allowed size of "
                f"{max_size // (1024 * 1024)} MB"
            )

        safe_name = _safe_filename(path.name)
        organization = _validate_scope(
            organization_id, "organization_id", "default"
        )
        resolved_namespace = _validate_scope(
            namespace, "namespace", "policy"
        )
        if resolved_namespace not in VALID_NAMESPACES:
            raise ValidationError(
                f"Unsupported namespace: {resolved_namespace}"
            )

        if document_id is None:
            document_hash = _file_sha256(path)
            document_id = f"doc_{document_hash[:32]}"
        else:
            document_id = str(document_id).strip()
            if (
                not document_id
                or len(document_id) > 128
                or not _SCOPE_PATTERN.fullmatch(document_id)
            ):
                raise ValidationError("Invalid document_id")

        extension = path.suffix.lower()

        if extension not in self.EXTENSION_BACKENDS:
            raise ValidationError(
                f"Unsupported file type: {extension or '[no extension]'}. "
                f"Supported types: "
                f"{', '.join(sorted(self.EXTENSION_BACKENDS))}"
            )

        backend_class = self.EXTENSION_BACKENDS[extension]

        if extension not in self.available_backends:
            raise RAGException(
                f"Required parser backend for '{extension}' is not "
                f"available. Install the corresponding dependency."
            )

        start_time = time.perf_counter()

        logger.info(
            "Parsing %s (%s) with %s",
            path.name,
            extension,
            backend_class.__name__,
        )

        try:
            parse_result = backend_class.parse(path)
        except Exception as error:
            logger.error(
                "Unhandled parser error for %s: %s",
                path.name,
                error,
                exc_info=True,
            )
            raise RAGException(
                f"Failed to parse {path.name}: {error}"
            ) from error

        if not isinstance(parse_result, dict):
            raise RAGException(
                f"Parser backend returned invalid result for "
                f"{path.name}"
            )

        text = parse_result.get("text", "")

        if text is None:
            text = ""

        if not isinstance(text, str):
            text = str(text)

        parse_result["text"] = text

        if len(text) > MAX_TEXT_LENGTH:
            raise RAGException(
                f"Parsed text from {safe_name} exceeds the maximum "
                f"supported text size"
            )

        if parse_result.get("error") and not text.strip():
            raise RAGException(
                f"Failed to parse {path.name}: "
                f"{parse_result['error']}"
            )

        if chunk and text.strip():
            parse_result["chunks"] = self.chunker.chunk(
                text=text,
                source=path.name,
                metadata={
                    "file_type": extension,
                    "pages": parse_result.get("pages", 0),
                    "ocr_used": parse_result.get("ocr_used", False),
                    "organization_id": organization,
                    "namespace": resolved_namespace,
                    "document_id": document_id,
                    "document_hash": document_id.removeprefix("doc_"),
                },
            )
        else:
            parse_result["chunks"] = []

        for index, item in enumerate(parse_result.get("chunks", [])):
            metadata = item.setdefault("metadata", {})
            metadata["chunk_index"] = index
            metadata["chunk_id"] = f"{document_id}_chunk_{index}"
            metadata["organization_id"] = organization
            metadata["namespace"] = resolved_namespace
            metadata["document_id"] = document_id
            metadata["document_hash"] = document_id.removeprefix("doc_")

        if extract_metadata:
            parse_result["metadata"] = self._extract_metadata(
                path,
                parse_result,
            )
            parse_result["metadata"].update({
                "organization_id": organization,
                "namespace": resolved_namespace,
                "document_id": document_id,
                "document_hash": document_id.removeprefix("doc_"),
            })

        elapsed_ms = (
            time.perf_counter() - start_time
        ) * 1000

        parse_result["parse_info"] = {
            "filename": path.name,
            "filepath": str(path.resolve()),
            "file_size_bytes": file_size,
            "file_type": extension,
            "parsed_at": _utc_now_iso(),
            "parse_time_ms": round(elapsed_ms, 2),
            "parser_backend": backend_class.__name__,
            "parse_method": parse_result.get(
                "parse_method",
                backend_class.__name__,
            ),
            "chunks_count": len(
                parse_result.get("chunks", [])
            ),
            "success": True,
        }

        logger.info(
            "Parsed %s: %s chars, %s chunks, %.0fms",
            path.name,
            len(text),
            len(parse_result["chunks"]),
            elapsed_ms,
        )

        return parse_result

    # -------------------------------------------------------------------------
    # Generic text parser
    # -------------------------------------------------------------------------

    def _parse_generic_text(
        self,
        file_path: Path,
    ) -> Dict[str, Any]:
        """
        Parse a plain-text file.

        This method is retained for backward compatibility but is intentionally
        NOT used as a fallback for unknown binary document formats.
        """
        result: Dict[str, Any] = {
            "text": "",
            "chunks": [],
            "parse_method": "generic_text",
        }

        try:
            with file_path.open(
                "r",
                encoding="utf-8",
                errors="replace",
            ) as file:
                result["text"] = file.read()

            result["chunks"] = self.chunker.chunk(
                text=result["text"],
                source=file_path.name,
            )

            logger.info(
                "Generic text parse: %s chars from %s",
                len(result["text"]),
                file_path.name,
            )

        except Exception as error:
            logger.error(
                "Generic text parse error: %s",
                error,
                exc_info=True,
            )
            result["error"] = str(error)

        return result

    # -------------------------------------------------------------------------
    # Metadata
    # -------------------------------------------------------------------------

    def _extract_metadata(
        self,
        file_path: Path,
        parse_result: Dict[str, Any],
    ) -> Dict[str, Any]:
        """Extract filesystem and parser metadata."""
        try:
            stat = file_path.stat()

            chunks = parse_result.get("chunks", [])

            chunk_sizes = [
                chunk.get("metadata", {}).get(
                    "char_length",
                    len(chunk.get("content", "")),
                )
                for chunk in chunks
            ]

            average_chunk_size = (
                sum(chunk_sizes) / len(chunk_sizes)
                if chunk_sizes
                else 0.0
            )

            return {
                "filename": file_path.name,
                "filepath": str(file_path),
                "file_size_bytes": stat.st_size,
                "file_size_mb": round(
                    stat.st_size / (1024 * 1024),
                    2,
                ),
                "created_at": datetime.fromtimestamp(
                    stat.st_ctime,
                    tz=timezone.utc,
                ).isoformat(),
                "modified_at": datetime.fromtimestamp(
                    stat.st_mtime,
                    tz=timezone.utc,
                ).isoformat(),
                "file_type": file_path.suffix.lower(),

                # Parser-specific information.
                "pages": parse_result.get("pages", 0),
                "paragraphs": parse_result.get("paragraphs", 0),
                "sheets": parse_result.get("sheets", 0),
                "slides": parse_result.get("slides", 0),
                "tables_count": parse_result.get("tables_count", 0),
                "ocr_used": parse_result.get("ocr_used", False),
                "ocr_confidence": parse_result.get(
                    "ocr_confidence",
                    0.0,
                ),

                # Chunking information.
                "chunks_count": len(chunks),
                "total_chars": len(
                    parse_result.get("text", "")
                ),
                "avg_chunk_size": round(
                    average_chunk_size,
                    2,
                ),
            }

        except Exception as error:
            logger.warning(
                "Metadata extraction error for %s: %s",
                file_path.name,
                error,
            )
            return {
                "filename": file_path.name,
                "error": str(error),
            }

    # -------------------------------------------------------------------------
    # Batch parsing
    # -------------------------------------------------------------------------

    def parse_batch(
        self,
        file_paths: List[Union[str, Path]],
        chunk: bool = True,
        on_progress: Optional[
            Callable[[int, int], None]
        ] = None,
        *,
        organization_id: Optional[str] = None,
        namespace: str = "policy",
    ) -> List[Dict[str, Any]]:
        """
        Parse multiple files concurrently.

        Results are returned in the SAME ORDER as file_paths, rather than
        completion order.
        """
        if not isinstance(file_paths, list):
            raise ValidationError(
                "file_paths must be a list"
            )

        if not file_paths:
            return []

        total = len(file_paths)
        results: List[Optional[Dict[str, Any]]] = [
            None
        ] * total

        completed = 0
        progress_lock = threading.Lock()

        def parse_one(
            index: int,
            file_path: Union[str, Path],
        ) -> Tuple[int, Dict[str, Any]]:
            path = _normalize_path(file_path)

            try:
                result = self.parse_file(
                    path,
                    chunk=chunk,
                    organization_id=organization_id,
                    namespace=namespace,
                )
                return index, result

            except Exception as error:
                logger.error(
                    "Failed to parse %s: %s",
                    path.name,
                    error,
                )

                return index, {
                    "error": str(error),
                    "filepath": str(path),
                    "filename": path.name,
                    "chunks": [],
                    "parse_info": {
                        "failed": True,
                        "success": False,
                        "parsed_at": _utc_now_iso(),
                    },
                }

        with ThreadPoolExecutor(
            max_workers=self.max_workers
        ) as executor:
            future_to_index = {
                executor.submit(
                    parse_one,
                    index,
                    file_path,
                ): index
                for index, file_path in enumerate(file_paths)
            }

            for future in as_completed(future_to_index):
                try:
                    index, result = future.result()
                    results[index] = result

                except Exception as error:
                    # Extremely defensive guard. parse_one already converts
                    # normal failures to result dictionaries.
                    index = future_to_index[future]
                    path = _normalize_path(file_paths[index])

                    results[index] = {
                        "error": str(error),
                        "filepath": str(path),
                        "filename": path.name,
                        "chunks": [],
                        "parse_info": {
                            "failed": True,
                            "success": False,
                            "parsed_at": _utc_now_iso(),
                        },
                    }

                with progress_lock:
                    completed += 1
                    current = completed

                if on_progress:
                    try:
                        on_progress(current, total)
                    except Exception as callback_error:
                        logger.warning(
                            "Progress callback failed: %s",
                            callback_error,
                        )

        final_results: List[Dict[str, Any]] = [
            result
            for result in results
            if result is not None
        ]

        successful = sum(
            1
            for result in final_results
            if "error" not in result
        )

        logger.info(
            "Batch parse complete: %s/%s successful",
            successful,
            total,
        )

        return final_results

    # -------------------------------------------------------------------------
    # Backend information
    # -------------------------------------------------------------------------

    def get_available_formats(self) -> List[str]:
        """Return formats whose parser dependencies are available."""
        return sorted(self.available_backends.keys())

    def get_backend_info(
        self,
        file_ext: str,
    ) -> Optional[Dict[str, Any]]:
        """Return information about a parser backend."""
        if not isinstance(file_ext, str):
            return None

        normalized_ext = file_ext.lower()

        if not normalized_ext.startswith("."):
            normalized_ext = "." + normalized_ext

        backend_class = self.EXTENSION_BACKENDS.get(
            normalized_ext
        )

        if backend_class is None:
            return None

        return {
            "extension": normalized_ext,
            "backend": backend_class.__name__,
            "available": normalized_ext in self.available_backends,
            "supports_ocr": backend_class in {
                PDFBackend,
                ImageBackend,
            },
        }

    def shutdown(self) -> None:
        """
        Release parser resources.

        Current backends do not maintain persistent resources, so this is
        intentionally lightweight.
        """
        logger.info(
            "MultimodalParser shutdown complete"
        )


# =============================================================================
# GLOBAL INSTANCE MANAGEMENT
# =============================================================================

_parser: Optional[MultimodalParser] = None
_parser_lock = threading.RLock()


def get_multimodal_parser(
    chunk_size: Optional[int] = None,
    chunk_overlap: Optional[int] = None,
    max_workers: int = DEFAULT_MAX_WORKERS,
) -> MultimodalParser:
    """
    Return the process-wide parser singleton.

    The first call determines the parser configuration. Subsequent calls
    return the existing instance.
    """
    global _parser

    with _parser_lock:
        if _parser is None:
            _parser = MultimodalParser(
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
                max_workers=max_workers,
            )

        return _parser


def reset_multimodal_parser() -> None:
    """Reset the global parser singleton, primarily for testing."""
    global _parser

    with _parser_lock:
        parser = _parser
        _parser = None

        if parser is not None:
            try:
                parser.shutdown()
            except Exception as error:
                logger.warning(
                    "Parser shutdown during reset failed: %s",
                    error,
                )


# =============================================================================
# CONVENIENCE FUNCTIONS
# =============================================================================

def parse_document(
    file_path: Union[str, Path],
    chunk: bool = True,
    *,
    organization_id: Optional[str] = None,
    namespace: str = "policy",
) -> Dict[str, Any]:
    """Parse a single document using the global parser."""
    return get_multimodal_parser().parse_file(
        file_path,
        chunk=chunk,
        organization_id=organization_id,
        namespace=namespace,
    )


def parse_documents_batch(
    file_paths: List[Union[str, Path]],
    chunk: bool = True,
    on_progress: Optional[
        Callable[[int, int], None]
    ] = None,
    *,
    organization_id: Optional[str] = None,
    namespace: str = "policy",
) -> List[Dict[str, Any]]:
    """Parse multiple documents using the global parser."""
    return get_multimodal_parser().parse_batch(
        file_paths,
        chunk=chunk,
        on_progress=on_progress,
    )


def get_parser_formats() -> List[str]:
    """Return currently available parser formats."""
    return get_multimodal_parser().get_available_formats()


# =============================================================================
# TEST / DEMO
# =============================================================================

def test_multimodal_parser() -> None:
    """Run a lightweight parser smoke test."""
    print("\nTesting Multimodal Parser")
    print("=" * 70)

    parser = get_multimodal_parser()

    print("Parser initialized")
    print(
        f"Available formats: "
        f"{parser.get_available_formats()}"
    )
    print(
        f"Config: chunk_size={parser.chunker.chunk_size}, "
        f"overlap={parser.chunker.chunk_overlap}, "
        f"workers={parser.max_workers}"
    )

    print("=" * 70)

    # -------------------------------------------------------------------------
    # Chunking smoke test
    # -------------------------------------------------------------------------

    print("\nTesting sentence-aware chunking")
    print("-" * 70)

    mock_policy = """
Employee Leave Policy.

1. Annual Leave: Employees are entitled to 20 days of paid leave per year.
Leave accrues monthly at the applicable rate. Unused leave may be carried
over according to company policy.

2. Sick Leave: Employees receive paid sick leave according to applicable
company policy. Sick leave may be used for personal illness or medical
appointments.

3. Maternity Leave: Eligible employees may receive maternity leave according
to applicable law and company policy.

4. Paternity Leave: Eligible employees may receive paternity leave according
to applicable law and company policy.

5. Leave Encashment: Unused leave may be eligible for encashment according
to applicable company policy.
""".strip()

    chunks = parser.chunker.chunk(
        text=mock_policy,
        source="test_policy.txt",
        metadata={"department": "HR"},
    )

    print(f"Input characters: {len(mock_policy)}")
    print(f"Generated chunks: {len(chunks)}")
    print(
        "Chunk sizes: "
        f"{[c['metadata']['char_length'] for c in chunks]}"
    )

    for index, chunk in enumerate(chunks[:3], start=1):
        print(
            f"\nChunk {index} "
            f"(ID={chunk['metadata']['chunk_id']}):"
        )
        print(
            f"  {chunk['content'][:120]}..."
        )

    # -------------------------------------------------------------------------
    # PDF smoke test
    # -------------------------------------------------------------------------

    pdf_dir = project_root / "data" / "PDFs"

    print("\nTesting PDF parsing")
    print("-" * 70)

    if pdf_dir.exists():
        pdf_files = sorted(
            pdf_dir.glob("*.pdf")
        )

        if pdf_files:
            print(
                f"Found {len(pdf_files)} PDF(s)"
            )

            for pdf_file in pdf_files[:2]:
                print(
                    f"\nParsing: {pdf_file.name}"
                )

                try:
                    result = parser.parse_file(
                        pdf_file
                    )

                    print(
                        f"  Pages: "
                        f"{result.get('pages', 0)}"
                    )
                    print(
                        f"  Chunks: "
                        f"{len(result.get('chunks', []))}"
                    )
                    print(
                        f"  Text: "
                        f"{len(result.get('text', '')):,} chars"
                    )
                    print(
                        f"  OCR: "
                        f"{result.get('ocr_used', False)}"
                    )
                    print(
                        f"  Parse time: "
                        f"{result.get('parse_info', {}).get('parse_time_ms', 0):.0f}ms"
                    )

                except Exception as error:
                    print(
                        f"  Error: {error}"
                    )
        else:
            print(
                f"No PDFs found in {pdf_dir}"
            )
    else:
        print(
            f"PDF directory not found: {pdf_dir}"
        )

    # -------------------------------------------------------------------------
    # Backend information
    # -------------------------------------------------------------------------

    print("\nBackend information")
    print("-" * 70)

    for extension in [
        ".pdf",
        ".docx",
        ".xlsx",
        ".pptx",
        ".png",
        ".doc",
        ".xls",
    ]:
        info = parser.get_backend_info(
            extension
        )

        if info:
            status = (
                "available"
                if info["available"]
                else "unavailable"
            )

            print(
                f"{extension:6s} -> "
                f"{info['backend']:15s} "
                f"({status}, OCR={info['supports_ocr']})"
            )

    print("\n" + "=" * 70)
    print("Multimodal parser smoke test complete")


if __name__ == "__main__":
    test_multimodal_parser()

