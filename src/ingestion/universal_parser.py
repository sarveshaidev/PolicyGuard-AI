
#!/usr/bin/env python3
"""
PolicyGuard AI - Universal Document Parser
===========================================

Unified parser for common HR document formats:
- PDF
- DOCX
- XLSX
- TXT

Features:
- PDF text extraction with OCR fallback
- DOCX paragraph/table extraction
- XLSX sheet/row extraction
- UTF-8 text parsing
- Sentence-aware configurable chunking
- Batch processing
- Thread-safe global singleton
- Metadata extraction
- Windows/Linux compatibility

Note:
Legacy .doc and .xls files are intentionally not advertised as supported.
python-docx does not parse legacy .doc files and openpyxl does not parse
legacy .xls files.
"""

import hashlib
import logging
import re
import sys
import threading
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

# ---------------------------------------------------------------------------
# Project path setup
# ---------------------------------------------------------------------------

current_file = Path(__file__).resolve()
project_root = current_file.parent.parent.parent

if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

from config.settings import settings
from src.core.exceptions import RAGException, ValidationError

logger = logging.getLogger(__name__)


# =============================================================================
# CONSTANTS / HELPERS
# =============================================================================

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt"}

DEFAULT_CHUNK_SIZE = 300
DEFAULT_CHUNK_OVERLAP = 50
MAX_CHUNK_SIZE = 1_000_000
MAX_WORKERS = 64
DEFAULT_MAX_FILE_SIZE_BYTES = 50 * 1024 * 1024
MAX_FILENAME_LENGTH = 255
MAX_TEXT_LENGTH = 10_000_000
VALID_NAMESPACES = {"policy", "talent"}
_SCOPE_PATTERN = re.compile(r"^[A-Za-z0-9_.:@-]{1,128}$")


def _utc_now_iso() -> str:
    """Return the current UTC time as ISO-8601."""
    return datetime.now(timezone.utc).isoformat()


def _setting_int(name: str, default: int) -> int:
    """Read an integer setting safely."""
    value = getattr(settings, name, default)

    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _normalize_path(file_path: Union[str, Path]) -> Path:
    """Validate and normalize a filesystem path."""
    if isinstance(file_path, Path):
        return file_path.expanduser()

    if not isinstance(file_path, str):
        raise ValidationError(
            "file_path must be a string or pathlib.Path"
        )

    if not file_path.strip():
        raise ValidationError(
            "file_path cannot be empty"
        )
    if "\x00" in file_path:
        raise ValidationError("file_path contains a null byte")

    return Path(file_path).expanduser()


def _validate_chunk_config(
    chunk_size: int,
    chunk_overlap: int,
) -> Tuple[int, int]:
    """Validate chunk size and overlap."""
    if isinstance(chunk_size, bool) or not isinstance(chunk_size, int):
        raise ValidationError(
            "chunk_size must be an integer"
        )

    if isinstance(chunk_overlap, bool) or not isinstance(chunk_overlap, int):
        raise ValidationError(
            "chunk_overlap must be an integer"
        )

    if chunk_size < 1 or chunk_size > MAX_CHUNK_SIZE:
        raise ValidationError(
            f"chunk_size must be between 1 and {MAX_CHUNK_SIZE}"
        )

    if chunk_overlap < 0:
        raise ValidationError(
            "chunk_overlap cannot be negative"
        )

    if chunk_overlap >= chunk_size:
        raise ValidationError(
            "chunk_overlap must be smaller than chunk_size"
        )

    return chunk_size, chunk_overlap



def _get_max_file_size_bytes() -> int:
    """Read the configured maximum file size with a safe fallback."""
    raw = getattr(settings, "MAX_UPLOAD_SIZE_MB", None)
    if raw is None:
        raw = getattr(settings, "MAX_FILE_SIZE_MB", None)
    try:
        mb = int(raw)
        if mb > 0:
            return mb * 1024 * 1024
    except (TypeError, ValueError):
        pass
    return DEFAULT_MAX_FILE_SIZE_BYTES


def _validate_scope(value: Optional[str], field_name: str, default: str) -> str:
    resolved = default if value is None else str(value).strip()
    if not resolved or not _SCOPE_PATTERN.fullmatch(resolved):
        raise ValidationError(f"Invalid {field_name}")
    return resolved


def _safe_filename(name: str) -> str:
    safe = Path(str(name)).name
    if not safe or safe in {".", ".."} or len(safe) > MAX_FILENAME_LENGTH:
        raise ValidationError("Invalid filename")
    return safe


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()

# =============================================================================
# UNIVERSAL PARSER
# =============================================================================

class UniversalParser:
    """
    Production-ready parser for common HR document formats.

    Supported:
    - PDF
    - DOCX
    - XLSX
    - TXT
    """

    SUPPORTED_EXTENSIONS = SUPPORTED_EXTENSIONS

    def __init__(
        self,
        chunk_size: Optional[int] = None,
        chunk_overlap: Optional[int] = None,
    ):
        configured_chunk_size = (
            _setting_int(
                "CHUNK_SIZE",
                DEFAULT_CHUNK_SIZE,
            )
            if chunk_size is None
            else chunk_size
        )

        configured_overlap = (
            _setting_int(
                "CHUNK_OVERLAP",
                DEFAULT_CHUNK_OVERLAP,
            )
            if chunk_overlap is None
            else chunk_overlap
        )

        (
            configured_chunk_size,
            configured_overlap,
        ) = _validate_chunk_config(
            configured_chunk_size,
            configured_overlap,
        )

        self.chunk_size = configured_chunk_size
        self.chunk_overlap = configured_overlap

        logger.info(
            "UniversalParser initialized: "
            "chunk_size=%s, overlap=%s",
            self.chunk_size,
            self.chunk_overlap,
        )

    # -------------------------------------------------------------------------
    # Main parsing
    # -------------------------------------------------------------------------

    def parse_file(
        self,
        file_path: Union[str, Path],
        *,
        organization_id: Optional[str] = None,
        namespace: str = "policy",
        document_id: Optional[str] = None,
    ) -> Dict[str, Any]:
        """
        Parse a supported file.

        Raises:
            FileNotFoundError
            ValidationError
            RAGException
        """
        path = _normalize_path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"File not found: {path}"
            )

        if not path.is_file():
            raise ValidationError(
                f"Path is not a file: {path}"
            )

        if path.is_symlink():
            raise ValidationError("Symbolic-link uploads are not allowed")

        try:
            file_size = path.stat().st_size
        except OSError as exc:
            raise ValidationError("Unable to inspect uploaded file") from exc

        if file_size <= 0:
            raise ValidationError("Uploaded file is empty")

        max_size = _get_max_file_size_bytes()
        if file_size > max_size:
            raise ValidationError(
                f"File exceeds the maximum allowed size of "
                f"{max_size // (1024 * 1024)} MB"
            )

        safe_name = _safe_filename(path.name)
        organization = _validate_scope(
            organization_id, "organization_id", "default"
        )
        resolved_namespace = _validate_scope(
            namespace, "namespace", "policy"
        )
        if resolved_namespace not in VALID_NAMESPACES:
            raise ValidationError(
                f"Unsupported namespace: {resolved_namespace}"
            )

        if document_id is None:
            document_hash = _file_sha256(path)
            document_id = f"doc_{document_hash[:32]}"
        else:
            document_id = str(document_id).strip()
            if (
                not document_id
                or len(document_id) > 128
                or not _SCOPE_PATTERN.fullmatch(document_id)
            ):
                raise ValidationError("Invalid document_id")

        extension = path.suffix.lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValidationError(
                f"Unsupported file extension: "
                f"{extension or '[no extension]'}. "
                f"Supported: "
                f"{', '.join(sorted(self.SUPPORTED_EXTENSIONS))}"
            )

        logger.info(
            "Parsing %s (%s)",
            path.name,
            extension,
        )

        start_time = time.perf_counter()

        try:
            if extension == ".pdf":
                result = self._parse_pdf(path)

            elif extension == ".docx":
                result = self._parse_docx(path)

            elif extension == ".xlsx":
                result = self._parse_xlsx(path)

            elif extension == ".txt":
                result = self._parse_txt(path)

            else:
                # Defensive guard in case the supported extension set changes.
                raise ValidationError(
                    f"No parser for extension: {extension}"
                )

        except (ValidationError, RAGException):
            raise

        except Exception as error:
            logger.error(
                "Unhandled parse error for %s: %s",
                path.name,
                error,
                exc_info=True,
            )
            raise RAGException(
                f"Failed to parse {path.name}: {error}"
            ) from error

        if not isinstance(result, dict):
            raise RAGException(
                f"Parser returned an invalid result for {path.name}"
            )

        text_value = result.get("text", "")
        if text_value is not None and len(str(text_value)) > MAX_TEXT_LENGTH:
            raise RAGException(
                f"Parsed text from {safe_name} exceeds the maximum "
                f"supported text size"
            )

        for index, item in enumerate(result.get("chunks", [])):
            metadata = item.setdefault("metadata", {})
            metadata["chunk_index"] = index
            metadata["chunk_id"] = f"{document_id}_chunk_{index}"
            metadata["organization_id"] = organization
            metadata["namespace"] = resolved_namespace
            metadata["document_id"] = document_id
            metadata["document_hash"] = document_id.removeprefix("doc_")

        # If a backend reports an error and produced no text, treat parsing as
        # failed rather than silently returning an empty document.
        if result.get("error") and not result.get("text", "").strip():
            raise RAGException(
                f"Failed to parse {path.name}: "
                f"{result['error']}"
            )

        elapsed_ms = (
            time.perf_counter() - start_time
        ) * 1000

        result["parse_info"] = {
            "filename": path.name,
            "filepath": str(path.resolve()),
            "file_size_bytes": file_size,
            "file_type": extension,
            "parsed_at": _utc_now_iso(),
            "parse_time_ms": round(elapsed_ms, 2),
            "chunks_count": len(
                result.get("chunks", [])
            ),
            "success": True,
        }

        result["metadata"] = self._extract_metadata(
            path,
            result,
        )
        result["metadata"].update({
            "organization_id": organization,
            "namespace": resolved_namespace,
            "document_id": document_id,
            "document_hash": document_id.removeprefix("doc_"),
        })

        logger.info(
            "Parsed %s: %s chunks, %.0fms",
            path.name,
            len(result.get("chunks", [])),
            elapsed_ms,
        )

        return result

    # -------------------------------------------------------------------------
    # PDF
    # -------------------------------------------------------------------------

    def _parse_pdf(
        self,
        file_path: Path,
    ) -> Dict[str, Any]:
        """Parse PDF with pypdf and OCR fallback."""
        result: Dict[str, Any] = {
            "text": "",
            "total_pages": 0,
            "ocr_used": False,
            "chunks": [],
            "parse_method": "pypdf",
        }

        try:
            from pypdf import PdfReader

            reader = PdfReader(str(file_path))
            result["total_pages"] = len(reader.pages)

            page_texts: List[Dict[str, Any]] = []

            for page_num, page in enumerate(
                reader.pages,
                start=1,
            ):
                try:
                    page_text = page.extract_text() or ""

                    if page_text.strip():
                        page_texts.append(
                            {
                                "page": page_num,
                                "content": page_text.strip(),
                            }
                        )

                except Exception as page_error:
                    logger.warning(
                        "Failed to extract PDF page %s: %s",
                        page_num,
                        page_error,
                    )

            result["text"] = "\n\n".join(
                f"[Page {item['page']}]\n{item['content']}"
                for item in page_texts
            )

            # Scanned/poorly extracted PDF.
            if len(result["text"].strip()) < 200:
                logger.info(
                    "PDF has limited extractable text (%s chars); "
                    "attempting OCR",
                    len(result["text"].strip()),
                )

                ocr_result = self._parse_pdf_ocr(
                    file_path,
                    existing_pages=result["total_pages"],
                )

                # Prefer OCR only when it actually produced text.
                if ocr_result.get("text", "").strip():
                    result = ocr_result
                else:
                    result["ocr_error"] = ocr_result.get(
                        "error"
                    )

        except ImportError:
            logger.warning(
                "pypdf is not installed; attempting OCR"
            )

            result = self._parse_pdf_ocr(
                file_path,
                existing_pages=0,
            )

        except Exception as error:
            logger.warning(
                "pypdf PDF parsing failed for %s: %s; "
                "attempting OCR",
                file_path.name,
                error,
            )

            ocr_result = self._parse_pdf_ocr(
                file_path,
                existing_pages=result["total_pages"],
            )

            if ocr_result.get("text", "").strip():
                result = ocr_result
            else:
                result["error"] = str(error)

        result["chunks"] = self._create_chunks(
            text=result.get("text", ""),
            source=file_path.name,
            page=1,
        )

        logger.info(
            "PDF parsed: %s pages, %s chars, OCR=%s",
            result.get("total_pages", 0),
            len(result.get("text", "")),
            result.get("ocr_used", False),
        )

        return result

    def _parse_pdf_ocr(
        self,
        file_path: Path,
        existing_pages: int = 0,
    ) -> Dict[str, Any]:
        """OCR fallback for scanned PDFs."""
        result: Dict[str, Any] = {
            "text": "",
            "total_pages": existing_pages,
            "ocr_used": True,
            "chunks": [],
            "parse_method": "ocr",
        }

        try:
            from src.vision.advanced_ocr import extract_text_from_pdf

            ocr_result = extract_text_from_pdf(file_path)

            if isinstance(ocr_result, dict):
                text = str(
                    ocr_result.get("text", "") or ""
                )

                if ocr_result.get("pages") is not None:
                    try:
                        result["total_pages"] = int(
                            ocr_result["pages"]
                        )
                    except (TypeError, ValueError):
                        pass

            else:
                text = str(
                    ocr_result or ""
                )

            result["text"] = text

            if text.strip():
                logger.info(
                    "PDF OCR successful: %s chars",
                    len(text),
                )
            else:
                logger.warning(
                    "PDF OCR returned no text"
                )

        except ImportError:
            result["error"] = (
                "OCR module is not available"
            )

        except Exception as error:
            logger.error(
                "PDF OCR error for %s: %s",
                file_path.name,
                error,
                exc_info=True,
            )
            result["error"] = str(error)

        return result

    # -------------------------------------------------------------------------
    # DOCX
    # -------------------------------------------------------------------------

    def _parse_docx(
        self,
        file_path: Path,
    ) -> Dict[str, Any]:
        """Parse DOCX paragraphs and tables."""
        result: Dict[str, Any] = {
            "text": "",
            "paragraphs": 0,
            "has_tables": False,
            "tables_count": 0,
            "chunks": [],
        }

        try:
            from docx import Document

            document = Document(str(file_path))

            paragraphs: List[str] = []

            for paragraph in document.paragraphs:
                text = paragraph.text.strip()

                if text:
                    paragraphs.append(text)

            result["paragraphs"] = len(paragraphs)
            result["text"] = "\n\n".join(paragraphs)

            if document.tables:
                result["has_tables"] = True
                result["tables_count"] = len(
                    document.tables
                )

                table_texts: List[str] = []

                for table_number, table in enumerate(
                    document.tables,
                    start=1,
                ):
                    rows: List[str] = []

                    for row in table.rows:
                        cells = [
                            cell.text.strip()
                            for cell in row.cells
                            if cell.text.strip()
                        ]

                        if cells:
                            rows.append(
                                " | ".join(cells)
                            )

                    if rows:
                        table_texts.append(
                            f"[Table {table_number}]\n"
                            + "\n".join(rows)
                        )

                if table_texts:
                    if result["text"]:
                        result["text"] += "\n\n"

                    result["text"] += "\n\n".join(
                        table_texts
                    )

        except ImportError:
            result["error"] = (
                "python-docx is not installed"
            )

        except Exception as error:
            logger.error(
                "DOCX parse error for %s: %s",
                file_path.name,
                error,
                exc_info=True,
            )
            result["error"] = str(error)

        result["chunks"] = self._create_chunks(
            text=result.get("text", ""),
            source=file_path.name,
            page=1,
        )

        logger.info(
            "DOCX parsed: %s paragraphs, %s tables, %s chars",
            result["paragraphs"],
            result["tables_count"],
            len(result["text"]),
        )

        return result

    # -------------------------------------------------------------------------
    # XLSX
    # -------------------------------------------------------------------------

    def _parse_xlsx(
        self,
        file_path: Path,
    ) -> Dict[str, Any]:
        """Parse XLSX worksheets into searchable text."""
        result: Dict[str, Any] = {
            "text": "",
            "sheets": 0,
            "rows": 0,
            "columns": 0,
            "chunks": [],
        }

        workbook = None

        try:
            from openpyxl import load_workbook

            workbook = load_workbook(
                str(file_path),
                read_only=True,
                data_only=True,
            )

            result["sheets"] = len(
                workbook.sheetnames
            )

            sheet_texts: List[str] = []
            total_rows = 0
            max_columns = 0

            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                rows_data: List[str] = []

                for row in sheet.iter_rows(
                    values_only=True
                ):
                    values: List[str] = []

                    for cell in row:
                        if cell is None:
                            values.append("")
                        else:
                            values.append(
                                str(cell).strip()
                            )

                    if not any(values):
                        continue

                    rows_data.append(
                        " | ".join(
                            value
                            for value in values
                            if value
                        )
                    )

                    total_rows += 1
                    max_columns = max(
                        max_columns,
                        len(values),
                    )

                if rows_data:
                    sheet_texts.append(
                        f"[Sheet: {sheet_name}]\n"
                        + "\n".join(rows_data)
                    )

            result["rows"] = total_rows
            result["columns"] = max_columns
            result["text"] = "\n\n".join(
                sheet_texts
            )

        except ImportError:
            result["error"] = (
                "openpyxl is not installed"
            )

        except Exception as error:
            logger.error(
                "XLSX parse error for %s: %s",
                file_path.name,
                error,
                exc_info=True,
            )
            result["error"] = str(error)

        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass

        result["chunks"] = self._create_chunks(
            text=result.get("text", ""),
            source=file_path.name,
            page=1,
        )

        logger.info(
            "XLSX parsed: %s sheets, %s rows, %s chars",
            result["sheets"],
            result["rows"],
            len(result["text"]),
        )

        return result

    # -------------------------------------------------------------------------
    # TXT
    # -------------------------------------------------------------------------

    def _parse_txt(
        self,
        file_path: Path,
    ) -> Dict[str, Any]:
        """Parse a UTF-8 text file."""
        result: Dict[str, Any] = {
            "text": "",
            "chunks": [],
            "parse_method": "text",
        }

        try:
            with file_path.open(
                "r",
                encoding="utf-8",
                errors="replace",
            ) as file:
                result["text"] = file.read()

        except Exception as error:
            logger.error(
                "TXT parse error for %s: %s",
                file_path.name,
                error,
                exc_info=True,
            )
            result["error"] = str(error)

        result["chunks"] = self._create_chunks(
            text=result.get("text", ""),
            source=file_path.name,
            page=1,
        )

        logger.info(
            "TXT parsed: %s chars",
            len(result["text"]),
        )

        return result

    # -------------------------------------------------------------------------
    # Chunking
    # -------------------------------------------------------------------------

    def _create_chunks(
        self,
        text: str,
        source: str,
        page: int = 1,
    ) -> List[Dict[str, Any]]:
        """
        Create overlapping sentence-aware chunks.

        Character positions refer to the original text before stripping.
        """
        if not isinstance(text, str):
            text = str(text or "")

        if not text.strip():
            return []

        chunks: List[Dict[str, Any]] = []

        text_length = len(text)
        start = 0
        chunk_index = 0

        while start < text_length:
            original_start = start
            end = min(
                start + self.chunk_size,
                text_length,
            )

            if end < text_length:
                boundary = self._find_boundary(
                    text,
                    start,
                    end,
                )

                if boundary > start:
                    end = boundary

            raw_text = text[start:end]
            chunk_text = raw_text.strip()

            if chunk_text:
                leading_trim = (
                    len(raw_text)
                    - len(raw_text.lstrip())
                )

                trailing_trim = (
                    len(raw_text)
                    - len(raw_text.rstrip())
                )

                char_start = start + leading_trim
                char_end = max(
                    char_start,
                    end - trailing_trim,
                )

                chunks.append(
                    {
                        "content": chunk_text,
                        "metadata": {
                            "source": source,
                            "page": page,
                            "chunk_index": chunk_index,
                            "char_start": char_start,
                            "char_end": char_end,
                            "char_length": len(chunk_text),
                        },
                        "score": 0.0,
                    }
                )

                chunk_index += 1

            # Always move forward.
            next_start = end - self.chunk_overlap

            if next_start <= original_start:
                next_start = original_start + 1

            if next_start >= text_length:
                break

            start = next_start

        average_size = (
            sum(
                chunk["metadata"]["char_length"]
                for chunk in chunks
            )
            / max(len(chunks), 1)
        )

        logger.debug(
            "Created %s chunks from '%s' "
            "(avg %.0f chars/chunk)",
            len(chunks),
            source,
            average_size,
        )

        return chunks

    def _find_boundary(
        self,
        text: str,
        start: int,
        end: int,
    ) -> int:
        """Find a sentence/paragraph boundary near chunk end."""
        search_text = text[start:end]

        # Search backwards through the last 150 characters.
        lower_bound = max(
            0,
            len(search_text) - 150,
        )

        # Prefer paragraph boundaries.
        paragraph_position = search_text.rfind(
            "\n\n",
            lower_bound,
        )

        if paragraph_position > len(search_text) // 2:
            return (
                start
                + paragraph_position
                + 2
            )

        # Then sentence endings.
        for index in range(
            len(search_text) - 1,
            lower_bound - 1,
            -1,
        ):
            if search_text[index] not in ".!?":
                continue

            previous_char = (
                search_text[index - 1]
                if index > 0
                else ""
            )

            next_char = (
                search_text[index + 1]
                if index + 1 < len(search_text)
                else ""
            )

            # Don't split decimal values such as 1.67.
            if (
                previous_char.isdigit()
                and next_char.isdigit()
            ):
                continue

            if index + 1 == len(search_text):
                return start + index + 1

            if next_char.isspace():
                return start + index + 1

        return end

    # -------------------------------------------------------------------------
    # Metadata
    # -------------------------------------------------------------------------

    def _extract_metadata(
        self,
        file_path: Path,
        parse_result: Dict[str, Any],
    ) -> Dict[str, Any]:
        """Extract filesystem and parsing metadata."""
        try:
            stat = file_path.stat()

            chunks = parse_result.get(
                "chunks",
                [],
            )

            chunk_sizes = [
                int(
                    chunk.get("metadata", {}).get(
                        "char_length",
                        len(
                            chunk.get(
                                "content",
                                "",
                            )
                        ),
                    )
                )
                for chunk in chunks
            ]

            average_chunk_size = (
                sum(chunk_sizes)
                / len(chunk_sizes)
                if chunk_sizes
                else 0.0
            )

            return {
                "filename": file_path.name,
                "filepath": str(file_path),
                "file_size_bytes": stat.st_size,
                "file_size_mb": round(
                    stat.st_size
                    / (1024 * 1024),
                    2,
                ),
                "created_at": datetime.fromtimestamp(
                    stat.st_ctime,
                    tz=timezone.utc,
                ).isoformat(),
                "modified_at": datetime.fromtimestamp(
                    stat.st_mtime,
                    tz=timezone.utc,
                ).isoformat(),
                "file_type": file_path.suffix.lower(),

                "pages": parse_result.get(
                    "total_pages",
                    0,
                ),
                "paragraphs": parse_result.get(
                    "paragraphs",
                    0,
                ),
                "sheets": parse_result.get(
                    "sheets",
                    0,
                ),
                "rows": parse_result.get(
                    "rows",
                    0,
                ),
                "ocr_used": parse_result.get(
                    "ocr_used",
                    False,
                ),
                "tables_count": parse_result.get(
                    "tables_count",
                    0,
                ),

                "chunks_count": len(chunks),
                "total_chars": len(
                    parse_result.get(
                        "text",
                        "",
                    )
                ),
                "avg_chunk_size": round(
                    average_chunk_size,
                    2,
                ),
            }

        except Exception as error:
            logger.warning(
                "Metadata extraction failed for %s: %s",
                file_path.name,
                error,
            )

            return {
                "filename": file_path.name,
                "error": str(error),
            }

    # -------------------------------------------------------------------------
    # Batch parsing
    # -------------------------------------------------------------------------

    def parse_batch(
        self,
        file_paths: List[Union[str, Path]],
        *,
        organization_id: Optional[str] = None,
        namespace: str = "policy",
    ) -> List[Dict[str, Any]]:
        """
        Parse multiple files.

        The method intentionally preserves the input order.
        """
        if not isinstance(file_paths, list):
            raise ValidationError(
                "file_paths must be a list"
            )

        if not file_paths:
            return []

        results: List[Dict[str, Any]] = []

        for file_path in file_paths:
            try:
                results.append(
                    self.parse_file(
                        file_path,
                        organization_id=organization_id,
                        namespace=namespace,
                    )
                )

            except Exception as error:
                try:
                    path = _normalize_path(
                        file_path
                    )
                    filepath = str(path)
                    filename = path.name
                except Exception:
                    filepath = str(file_path)
                    filename = str(file_path)

                logger.error(
                    "Failed to parse %s: %s",
                    filename,
                    error,
                )

                results.append(
                    {
                        "error": str(error),
                        "filepath": filepath,
                        "filename": filename,
                        "chunks": [],
                        "parse_info": {
                            "success": False,
                            "failed": True,
                            "parsed_at": _utc_now_iso(),
                        },
                    }
                )

        successful = sum(
            1
            for result in results
            if "error" not in result
        )

        logger.info(
            "Batch parse complete: %s/%s successful",
            successful,
            len(results),
        )

        return results


# =============================================================================
# GLOBAL INSTANCE MANAGEMENT
# =============================================================================

_parser: Optional[UniversalParser] = None
_parser_lock = threading.RLock()


def get_universal_parser(
    chunk_size: Optional[int] = None,
    chunk_overlap: Optional[int] = None,
) -> UniversalParser:
    """Get or create the global UniversalParser instance."""
    global _parser

    with _parser_lock:
        if _parser is None:
            _parser = UniversalParser(
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
            )

        return _parser


def reset_universal_parser() -> None:
    """Reset the global parser singleton, primarily for tests."""
    global _parser

    with _parser_lock:
        _parser = None


# =============================================================================
# CONVENIENCE FUNCTIONS
# =============================================================================

def parse_document(
    file_path: Union[str, Path],
    *,
    organization_id: Optional[str] = None,
    namespace: str = "policy",
) -> Dict[str, Any]:
    """Quickly parse one document."""
    return get_universal_parser().parse_file(
        file_path,
        organization_id=organization_id,
        namespace=namespace,
    )


def parse_documents_batch(
    file_paths: List[Union[str, Path]],
    *,
    organization_id: Optional[str] = None,
    namespace: str = "policy",
) -> List[Dict[str, Any]]:
    """Quickly parse multiple documents."""
    return get_universal_parser().parse_batch(
        file_paths,
        organization_id=organization_id,
        namespace=namespace,
    )


# =============================================================================
# TEST / DEMO
# =============================================================================

def test_universal_parser() -> None:
    """Run a basic UniversalParser smoke test."""
    print("\nTesting Universal Parser")
    print("=" * 70)

    parser = get_universal_parser()

    print("Parser initialized")
    print(
        "Supported formats: "
        f"{sorted(parser.SUPPORTED_EXTENSIONS)}"
    )
    print(
        "Config: "
        f"chunk_size={parser.chunk_size}, "
        f"overlap={parser.chunk_overlap}"
    )

    print("=" * 70)

    # -------------------------------------------------------------------------
    # Chunking test
    # -------------------------------------------------------------------------

    print("\nTesting chunking")
    print("-" * 70)

    mock_text = (
        "Sentence one. "
        "Sentence two! "
        "Sentence three? "
    ) * 20

    chunks = parser._create_chunks(
        text=mock_text,
        source="test.txt",
        page=1,
    )

    print(
        f"Input: {len(mock_text)} characters"
    )
    print(
        f"Output: {len(chunks)} chunks"
    )

    if chunks:
        print(
            "First chunk size: "
            f"{chunks[0]['metadata']['char_length']}"
        )

    # -------------------------------------------------------------------------
    # PDF test
    # -------------------------------------------------------------------------

    print("\nTesting PDF parsing")
    print("-" * 70)

    pdf_dir = project_root / "data" / "PDFs"

    if not pdf_dir.exists():
        print(
            f"PDF directory not found: {pdf_dir}"
        )
    else:
        pdf_files = sorted(
            pdf_dir.glob("*.pdf")
        )

        if not pdf_files:
            print(
                f"No PDFs found in {pdf_dir}"
            )
        else:
            print(
                f"Found {len(pdf_files)} PDF(s)"
            )

            for pdf_file in pdf_files[:2]:
                print(
                    f"\nParsing: {pdf_file.name}"
                )

                try:
                    result = parser.parse_file(
                        pdf_file
                    )

                    print(
                        f"  Pages: "
                        f"{result.get('total_pages', 0)}"
                    )
                    print(
                        f"  Chunks: "
                        f"{len(result.get('chunks', []))}"
                    )
                    print(
                        f"  Text: "
                        f"{len(result.get('text', '')):,} chars"
                    )
                    print(
                        f"  OCR: "
                        f"{result.get('ocr_used', False)}"
                    )
                    print(
                        "  Parse time: "
                        f"{result.get('parse_info', {}).get('parse_time_ms', 0):.0f}ms"
                    )

                except Exception as error:
                    print(
                        f"  Error: {error}"
                    )

    print("\n" + "=" * 70)
    print("Universal parser smoke test complete")


if __name__ == "__main__":
    test_universal_parser()

